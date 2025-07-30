from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
"""
Script to extract the markdown file and section containing a value marked by a localization id (e.g., @@loc_123@@).
Given a localization marker, a markdown directory, a demarked directory, and a localization map JSON, this script finds and highlights all sections containing the localization value.
"""

import re
import json
import html
from pathlib import Path
from typing import List, Tuple, Union

def normalize_text(text: str) -> str:
    """Normalize text for flexible searching: lowercase, replace <br> and <br/> with space, remove extra whitespace, unescape HTML."""
    text = html.unescape(text)
    text = text.replace('<br>', ' ').replace('<br/>', ' ')
    text = text.replace('\n', ' ')
    text = ' '.join(text.lower().split())
    return text

def highlight_all(text: str, pattern: str, norm_pattern: str) -> str:
    """Highlight all matches of pattern and normalized pattern in the text."""
    # Remove any existing highlight spans to prevent nesting
    clean_text = re.sub(r"<span style='background-color: yellow'>(.*?)</span>", r"\1", text)
    def replacer(match):
        # Escape parentheses to prevent markdown link formatting
        highlighted_text = match.group(0).replace('(', '\\(').replace(')', '\\)')
        return f"<span style='background-color: yellow'>{highlighted_text}</span>"
    highlighted = re.sub(pattern, replacer, clean_text, flags=re.IGNORECASE)
    if pattern != norm_pattern:
        highlighted = re.sub(norm_pattern, replacer, normalize_text(highlighted), flags=re.IGNORECASE)
    return highlighted

def write_docx_with_highlight(doc, header: str, body: str, match: str):
    """Append the header and body to a .docx Document, highlighting the match in yellow."""
    doc.add_heading(header.strip('#').strip(), level=2)
    lines = body.split('\n') if '\\n' in body else body.split('\n')
    in_table = False
    table_lines = []
    for line in lines:
        if line.strip().startswith('|') and line.strip().endswith('|'):
            in_table = True
            table_lines.append(line)
        else:
            if in_table and table_lines:
                # Process table
                add_markdown_table_to_docx(doc, table_lines, match)
                table_lines = []
                in_table = False
            if line.strip():
                add_paragraph_with_highlight(doc, line, match)
    if in_table and table_lines:
        add_markdown_table_to_docx(doc, table_lines, match)

def add_paragraph_with_highlight(doc, text, match):
    para = doc.add_paragraph()
    start = 0
    match_lower = match.lower()
    text_lower = text.lower()
    while True:
        idx = text_lower.find(match_lower, start)
        if idx == -1:
            para.add_run(text[start:])
            break
        para.add_run(text[start:idx])
        run = para.add_run(text[idx:idx+len(match)])
        # Highlight in yellow
        highlight = OxmlElement('w:highlight')
        highlight.set(qn('w:val'), 'yellow')
        run.font.color.rgb = RGBColor(0,0,0)
        rPr = run._element.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            run._element.insert(0, rPr)
        rPr.append(highlight)
        start = idx + len(match)

def add_markdown_table_to_docx(doc, table_lines, match):
    # Parse markdown table
    rows = [
        [cell.strip() for cell in row.strip().strip('|').split('|')] for row in table_lines if row.strip()
    ]
    if not rows:
        return
    table = doc.add_table(rows=0, cols=len(rows[0]))
    for row in rows:
        row_cells = table.add_row().cells
        for i, cell in enumerate(row):
            add_text_with_highlight_to_cell(row_cells[i], cell, match)

def add_text_with_highlight_to_cell(cell, text, match):
    para = cell.paragraphs[0]
    start = 0
    match_lower = match.lower()
    text_lower = text.lower()
    while True:
        idx = text_lower.find(match_lower, start)
        if idx == -1:
            para.add_run(text[start:])
            break
        para.add_run(text[start:idx])
        run = para.add_run(text[idx:idx+len(match)])
        # Highlight in yellow
        highlight = OxmlElement('w:highlight')
        highlight.set(qn('w:val'), 'yellow')
        run.font.color.rgb = RGBColor(0,0,0)
        rPr = run._element.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            run._element.insert(0, rPr)
        rPr.append(highlight)
        start = idx + len(match)

def find_section_with_marker(
    loc_marker: str,
    md_dir: Union[str, Path],
    demarked_dir: Union[str, Path],
    json_path: Union[str, Path]
) -> List[Tuple[str, str]]:
    md_dir = Path(md_dir)
    demarked_dir = Path(demarked_dir)
    json_path = Path(json_path)
    with open(json_path, 'r', encoding='utf-8') as jf:
        loc_map = json.load(jf)
    file_name = None
    marker_stripped = loc_marker.strip('@')
    localization_value = None
    for k, v in loc_map.items():
        if k == marker_stripped or ('@@' + k + '@@') == loc_marker:
            file_name = v.get('file')
            localization_value = v.get('localization', None)
            break
    if not file_name:
        for k, v in loc_map.items():
            if loc_marker in v.get('localization', ''):
                file_name = v.get('file')
                localization_value = v.get('localization', None)
                break
    if not file_name:
        print(f"[DEBUG] Marker {loc_marker} not found in localization map.")
        return []
    orig_md_file = md_dir / Path(file_name).stem / (Path(file_name).stem + '.md')
    print(f"[DEBUG] Searching file: {orig_md_file}")
    if not orig_md_file.exists():
        print(f"[DEBUG] Markdown file does not exist: {orig_md_file}")
        return []
    with open(orig_md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    sections = re.split(r'(^#+ .*$)', content, flags=re.MULTILINE)
    if not localization_value:
        print(f"[DEBUG] No localization value found for marker {loc_marker}")
        return []
    search_term = normalize_text(localization_value)
    highlight_pattern = re.escape(localization_value)
    norm_pattern = re.escape(search_term)
    print(f"[DEBUG] Search terms: {localization_value}")
    all_md = []
    output_file = demarked_dir / (Path(file_name).stem + '.out.md')
    docx_output = demarked_dir / (Path(file_name).stem + '.docx')
    docx_doc = Document()
    found_any = False
    matches = []
    for i in range(1, len(sections), 2):
        header = sections[i].rstrip('\n')
        body = sections[i+1] if i+1 < len(sections) else ''
        norm_body = normalize_text(body)
        found = search_term in norm_body
        print(f"[DEBUG] Section header: {header[:60]}... | Found: {found}")
        if found:
            found_any = True
            highlighted_body = highlight_all(body, highlight_pattern, norm_pattern)
            all_md.append(f"{header}\n\n{highlighted_body}\n{'-'*40}\n")
            # Add to docx_doc (accumulate all matches)
            write_docx_with_highlight(docx_doc, header, body, localization_value)
            matches.append((str(orig_md_file), header + '\n' + highlighted_body))
    if found_any:
        with open(output_file, 'w', encoding='utf-8') as out_f:
            out_f.writelines(all_md)
        docx_doc.save(str(docx_output))
    else:
        with open(output_file, 'w', encoding='utf-8') as out_f:
            out_f.write("Marker not found in any section.\n")
    print(f"[DEBUG] Total matches found: {len(matches)}")
    return matches

if __name__ == "__main__":
    import sys
    default_loc_marker = "@@loc_123@@"
    default_md_dir = "../main_storage/exemple2/markdown"
    default_demarked_dir = "../main_storage/exemple2/debugsections"
    default_json_path = "../main_storage/exemple2/localizations/localizations_map.json"
    if len(sys.argv) >= 5:
        loc_marker = sys.argv[1]
        md_dir = sys.argv[2]
        demarked_dir = sys.argv[3]
        json_path = sys.argv[4]
    else:
        print("[DEBUG] Using default arguments:")
        print(f"  loc_marker: {default_loc_marker}")
        print(f"  md_dir: {default_md_dir}")
        print(f"  demarked_dir: {default_demarked_dir}")
        print(f"  json_path: {default_json_path}")
        loc_marker = default_loc_marker
        md_dir = default_md_dir
        demarked_dir = default_demarked_dir
        json_path = default_json_path
    results = find_section_with_marker(loc_marker, md_dir, demarked_dir, json_path)
    if results:
        for idx, (file_path, section) in enumerate(results, 1):
            print(f"Match {idx} - File: {file_path}\nSection:\n{section}\n{'-'*40}")
